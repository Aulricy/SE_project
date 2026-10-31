from pathlib import Path
import sys,json,re
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.section import WD_SECTION_START,WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
sys.stdout.reconfigure(encoding='utf-8')
ROOT=Path(__file__).resolve().parents[1]
WORK=Path(__file__).resolve().parent
FIG=WORK/'插图'
OUT=ROOT/'实验报告.docx'
doc=Document()
doc.core_properties.title='二手交易平台软件设计与建模实验报告'
doc.core_properties.subject='软件工程实验二'
doc.core_properties.author='胥松恒'
from datetime import datetime,timezone
doc.core_properties.created=datetime(2026,10,31,tzinfo=timezone.utc)
doc.core_properties.modified=datetime.now(timezone.utc)
normal=doc.styles['Normal'];normal.font.name='Microsoft YaHei';normal.font.size=Pt(11)
normal.element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
normal.paragraph_format.line_spacing=Pt(16.5)
normal.paragraph_format.space_after=Pt(7)
for n,size in [('Title',25),('Subtitle',15),('Heading 1',18),('Heading 2',13),('Heading 3',11)]:
 s=doc.styles[n];s.font.name='Microsoft YaHei';s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0);s.font.italic=False;s.element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
 s.paragraph_format.space_before=Pt(9);s.paragraph_format.space_after=Pt(7)
 for border in list(s.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
 s.paragraph_format.keep_with_next=True
 s.paragraph_format.line_spacing=Pt(size*1.3)
for sec in doc.sections:
 sec.page_width=Inches(8.5);sec.page_height=Inches(11)
 sec.top_margin=sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.7)
 sec.header_distance=sec.footer_distance=Inches(.25)
header=doc.sections[0].header.paragraphs[0];header.text='软件工程实验二  二手交易平台';header.style=doc.styles['Normal']
header.runs[0].font.size=Pt(8);header.paragraph_format.space_after=Pt(0)
f=doc.sections[0].footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=f.add_run('');r.font.size=Pt(8)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
orientation='portrait';page_manifest=[]
def page(title,landscape=False):
 global orientation
 target='landscape' if landscape else 'portrait'
 previous_orientation=orientation
 if page_manifest:
  if target!=orientation:
   s=doc.add_section(WD_SECTION_START.NEW_PAGE)
   s.orientation=WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
   s.page_width=Inches(11 if landscape else 8.5);s.page_height=Inches(8.5 if landscape else 11)
   s.top_margin=s.bottom_margin=Inches(.6 if landscape else .65)
  else:pass
 orientation=target;page_manifest.append({'title':title,'landscape':landscape})
 if title:
  h=doc.add_heading(title,level=1)
  h.paragraph_format.page_break_before=bool(page_manifest[:-1]) and target==previous_orientation
def para(t):
 return doc.add_paragraph(t)
def sub(t):doc.add_heading(t,level=2)
def table(head,rows,widths=None,size=10):
 t=doc.add_table(rows=1, cols=len(head));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 if widths:
  for c,w in zip(t.columns,widths):c.width=Inches(w)
 for i,x in enumerate(head):t.rows[0].cells[i].text=x
 for row in rows:
  cells=t.add_row().cells
  for c,v in zip(cells,row):c.text=str(v)
 for ri,row in enumerate(t.rows):
  trpr=row._tr.get_or_add_trPr();ns=OxmlElement('w:cantSplit');trpr.append(ns)
  if ri==0:
   rep=OxmlElement('w:tblHeader');trpr.append(rep)
  for ci,c in enumerate(row.cells):
   if widths:c.width=Inches(widths[ci])
   c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   tcpr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
   tcpr.append(b)
   margins=OxmlElement('w:tcMar')
   for edge in ['top','bottom','left','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:w'),'70');e.set(qn('w:type'),'dxa');margins.append(e)
   tcpr.append(margins)
   fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'E7EDF3' if ri==0 else ('F7F9FB' if ri%2==0 else 'FFFFFF'));tcpr.append(fill)
   for p in c.paragraphs:
    p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=Pt(13.5)
    for r in p.runs:r.font.size=Pt(size);r.bold=(ri==0)
 spacer=doc.add_paragraph();spacer.paragraph_format.space_after=Pt(4);spacer.paragraph_format.line_spacing=1
 spacer.add_run(' ').font.size=Pt(1)
 return t
def guide(text):
 p=doc.add_paragraph(text);p.paragraph_format.line_spacing=Pt(12);p.paragraph_format.space_after=Pt(5)
 for run in p.runs:run.font.size=Pt(9.5)
def cellpic(cell,path,maxw,maxh):
 from PIL import Image
 with Image.open(path) as im:iw,ih=im.size
 width=min(maxw,maxh*iw/ih)
 p=cell.paragraphs[0];p.paragraph_format.line_spacing=1.0;p.paragraph_format.space_after=Pt(0);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 p.add_run().add_picture(str(path),width=Inches(width))

def pic(path,caption,width=7.1,maxheight=8.1):
 from PIL import Image
 with Image.open(path) as im:w,h=im.size
 width=min(width,maxheight*w/h)
 p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.0
 p.add_run().add_picture(str(path),width=Inches(width))
 p=doc.add_paragraph(caption);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(5)
 for r in p.runs:r.font.size=Pt(9)

page('')
p=doc.add_paragraph('二手交易平台\n软件设计与建模实验报告',style='Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before=Pt(55);p.paragraph_format.space_after=Pt(20)
p=doc.add_paragraph('软件工程 2026 秋  实验二',style='Subtitle');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
for t in ['姓名  胥松恒','学号  241220054','报告日期  2026 年 10 月 31 日']:
 p=para(t);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(13)
sub('报告内容')
para('本实验基于实验一的二手交易平台需求，设计用户与权限、商品管理、预约与交易、交易统计、操作日志五个模块。通过类图、组件图、活动图、时序图说明数据结构与业务协作，并以离线界面原型展示主要操作和异常状态。')
para('设计采用同一 Web 应用和共享关系数据库。预约、确认、取消、完成以及屏蔽引起的联动更新以事务为边界；交易在线下付款和交付，系统只记录预约与完成结果。')
sub('阅读顺序')
para('一 项目与功能  →  二 模块划分  →  三 UML 设计  →  四 界面设计  →  五 验证与需求对应  →  六 AI 使用过程')

page('一 项目与主要功能')
sub('1 项目目标与依据')
para('项目为二手交易平台，支持卖家发布闲置物品、买家预约商品、双方在线下完成交易，以及管理员审核和处理违规商品。设计依据为实验一仓库 Aulricy/SE_project 的 lab1-final 标签下 requirements.md，共 22 条需求；本实验不扩大既定业务范围。')
sub('2 参与者与权限')
table(['参与者','可执行操作','权限边界'],[
 ['游客','浏览和组合查询已发布商品，查看公开描述','预约及个人记录需要登录'],
 ['普通用户','同一账号既可发布商品，也可预约他人商品','仅管理自己的商品和相关预约'],
 ['管理员','审核、屏蔽、恢复商品，查询全平台统计和日志','不能代买家完成交易，不能改写完成历史'],
 ['首个管理员','具有管理员权限，并批准设置新增管理员账号','以预置账号编号识别，其他管理员不能继续增设']
],[1.05,3.0,3.05])
sub('3 主要功能与业务范围')
para('商品按数码、图书、生活用品、服饰、其他五类管理。一条记录代表一件物品；名称、分类、价格、描述是提交审核的必填内容，图片最多一张且可不上传。草稿可以不完整，成色、联系方式和交付说明统一写在描述中。')
para('买家只能预约他人已发布商品。同一商品最多有一条待确认或已确认预约。卖家确认后，买家在线下交付完成后标记完成；双方也可在待确认或已确认阶段填写原因取消。历史预约和预约价格始终保留。')
para('系统不提供在线支付、物流、聊天、推荐、议价、多库存、多图、自动过期或候补排队。管理员后台只是既有模块的管理入口，不另设第六个业务模块。')

page('二 模块划分与协作')
sub('1 五个模块的职责')
table(['模块','具体功能','主要协作对象'],[
 ['用户与权限','普通注册、统一登录、角色与归属校验、首个管理员批准设置账号','供所有受限业务调用'],
 ['商品管理','草稿、提交与撤回审核、审核与驳回、公开查询、下架、退回草稿、屏蔽与恢复','用户与权限、预约与交易、日志'],
 ['预约与交易','建立预约、卖家确认、双方取消、买家完成、买入卖出记录、重复请求识别','商品、用户与权限、日志'],
 ['交易统计','按首次完成的北京时间日期汇总完成笔数和预约价格合计','读取已完成预约，校验管理员'],
 ['操作日志','业务内部写入成功关键操作，管理员只读组合查询','业务模块写入，查询时鉴权']
],[1.05,3.95,2.1])
sub('2 模块之间如何配合')
para('卖家先经用户与权限模块确认身份，再由商品模块保存草稿和提交审核。管理员审核通过后，商品才进入公开查询；审核、驳回、撤回等关键成功操作由日志模块记录。将审核放在商品模块内，可集中维护商品生命周期，避免多个模块分别解释同一商品状态。')
para('买家预约时，预约模块读取并锁定商品，检查身份、归属、状态与版本，再联动保存预约和商品状态。卖家确认、双方取消、买家完成均沿用这一边界。商品屏蔽由商品模块发起，并调用预约模块内部的取消能力，保证屏蔽与有效预约取消同时生效。')
para('交易统计只读取已完成预约，不参与修改交易。日志查询也不修改历史。五个模块置于同一 Web 应用，使用共享数据库，便于用单个业务事务处理跨模块更新；当前规模无需引入独立微服务及分布式事务。')

page('三 UML 设计与关键对象')
para('本报告绘制四种 UML：类图和组件图属于结构性图，活动图属于行为性图，时序图属于交互性图。活动图按业务拆成商品发布审核、预约交易两张，合计五幅完整图；原始 PlantUML 源文件和 PNG 随成果保存。')
sub('1 类图中的对象')
table(['对象','主要属性','设计理由'],[
 ['User','id、username、passwordHash、role、createdAt','买卖共用账号；密码保存完整带盐校验信息'],
 ['Product','名称、分类、价格、描述、可选图片、status、各时间、最近驳回及屏蔽原因、version','允许不完整草稿；状态和版本控制后续操作'],
 ['Reservation','预约价格、status、创建/确认/取消/完成时间、取消原因、version','价格快照不随商品改价而变化，保留每次预约历史'],
 ['OperationLog','操作类型、时间、商品及预约前后状态、原因','支持追溯；完整历史不由商品最近原因替代'],
 ['操作请求记录','请求标识、操作类型、请求指纹、创建时间','识别成功请求重发，关联原预约和操作者'],
 ['StatisticsResult','日期范围、完成笔数、预约价格合计','查询结果对象，按需计算，不单独作为业务实体保存']
],[1.35,3.2,2.55],9.5)
sub('2 关联与职责')
para('一个用户可拥有多件商品，也可作为买家关联多条预约。一件商品可关联多条历史预约，但有效预约最多一条。预约通过商品找到卖家，避免重复存储容易不一致的卖家字段。日志关联操作者和商品，预约关联可选；请求记录则关联操作者和原预约。')
para('采用普通关联保留历史，不以组合关系暗示删除商品会级联删除预约或日志。业务操作集中在五个服务中；枚举属性说明角色、分类和状态，不再为每个枚举重复绘制连线。')

page('长图阅读导航',True)
para('以下缩略全貌用于定位分区，节点文字和连线请按 A、B、C 标记查看后面的放大图。三幅图分别保持原来的整体结构。')
t=table(['类图','商品发布与审核活动图','预约与交易活动图'],[['','',''],['A 业务实体 → 图1a\nB 服务协作 → 图1b\nC 枚举 → 状态与一致性约束','A 草稿与审核 → 图3a\nB 当前状态出口 → 图3b','A 建立预约及待确认 → 图4a\nB 交付与完成处理 → 图4b']],[3.1,3.2,3.2],10)
for cell,name in zip(t.rows[1].cells,['导航_类图','导航_发布审核','导航_预约交易']):cellpic(cell,FIG/(name+'.png'),2.95,3.4)
guide('阅读顺序：先确认整体位置，再沿相同标记阅读分区。时序图按调用顺序分上下两部分，续图顶部重复参与者名称。')

page('类图中的业务实体与关联')
guide('导航 A  业务实体与关联；后续 B 为服务协作，C 为枚举与状态约束。')
pic(FIG/'类图_实体.png','图 1a  类图的业务实体分区  完整类图见 UML/类图.png',maxheight=7.8)
para('端点数字表示多重性。卖家、买家、取消人及操作者通过对象关联表达；同一商品允许多条历史预约。有效预约唯一性属于跨记录约束，需要在数据库与事务中共同保证。')

page('类图中的服务协作')
guide('导航 B  承接实体分区，展示五个服务的调用依赖；省略的完整参数见源码注释。')
pic(FIG/'类图_服务.png','图 1b  类图的服务及查询结果分区',maxheight=7.8)
para('省略号只用于控制图的宽度，完整方法参数保存在 .puml 注释中。写操作明确接收操作者与预期版本；预约操作还携带请求标识。屏蔽联动取消和日志写入仅供模块内部使用。')

page('类图的状态与一致性约束')
sub('3 状态对应关系')
table(['场景','商品状态','预约状态'],[
 ['草稿、待审核、审核通过','DRAFT → PENDING_REVIEW → PUBLISHED','尚未建立预约'],
 ['建立预约','PUBLISHED → RESERVED','PENDING_CONFIRMATION'],
 ['卖家确认','RESERVED → CONFIRMED','CONFIRMED'],
 ['双方普通取消或卖家拒绝','恢复 PUBLISHED','CANCELLED，保留原因、取消人和时间'],
 ['买家完成已确认交易','CONFIRMED → COMPLETED','COMPLETED，首次完成时间只写一次'],
 ['管理员屏蔽未完成商品','BLOCKED','存在有效预约时联动设为 CANCELLED'],
 ['管理员恢复已屏蔽商品','BLOCKED → DRAFT','旧预约仍为 CANCELLED'],
 ['卖家下架或返回草稿','无有效预约时进入 OFF_SHELF 或 DRAFT','历史记录保持不变']
],[2.05,2.25,2.8],10)
para('商品分类固定为 DIGITAL、BOOKS、DAILY_NECESSITIES、CLOTHING、OTHER；用户角色仅 MEMBER、ADMIN。首个管理员的额外能力以只读配置 initialAdminId 校验，无需增加第三种账号角色。已完成交易不允许取消、重发或被屏蔽处理改写。')
sub('4 类图的简化与修订')
para('最终类图将实体、服务、枚举分区展示，缩短方法展示并将完整签名留在源码中。草稿的新建与修改分别使用 createDraft、updateDraft，避免把创建时由系统设置的编号、卖家和状态作为用户输入。商品写操作统一校验 expectedVersion，防止旧页面覆盖新状态。')
para('新增管理员由 createAdmin 执行，并先调用 requireInitialAdmin，明确区分“具有管理员角色”与“首个管理员”。这些调整保留必要约束，同时避免为统计结果、联系方式、申请队列等增加无需求依据的实体。')

page('组件图与架构选择')
pic(ROOT.parent/'UML'/'组件图.png','图 2  二手交易平台组件图',maxheight=7.6)
para('界面通过 HTTP 业务接口调用同一应用内的五个组件，不直接访问数据库。关系库保存用户、商品、预约、请求记录和日志；图片以本地文件保存，数据库记录路径。商品、预约、请求记录和关键日志应在同一事务中提交。')
page('商品发布与审核活动图上半部分',True)
guide('导航 A  从保存草稿进入提交、撤回、审核、驳回或屏蔽；后续状态出口见 B。')
pic(FIG/'发布活动_上.png','图 3a  草稿保存与提交  撤回 审核 驳回 屏蔽分支',width=9.55,maxheight=5.25)
para('草稿保存与提交审核分开；提交时再检查必填内容、价格、所属用户和版本。两部分展示同一幅活动图，中间保留重叠区域用于衔接。')
page('商品发布与审核活动图下半部分',True)
guide('导航 B  承接 A 的处理结果，重新判断当前状态；右侧回线返回 A 的草稿编辑。')
pic(FIG/'发布活动_下.png','图 3b  处理失败与当前状态出口  承接图 3a',width=9.55,maxheight=5.25)
para('最终按当前真实状态判断出口：只有已发布表示发布成功；已屏蔽终止本次流程，草稿可继续修改，仍待审核或其他状态则刷新。冲突失败不自动覆盖数据。')
page('预约与交易活动图上半部分',True)
guide('导航 A  建立一次预约，处理待确认阶段；确认成功后沿下接线进入 B。')
pic(FIG/'交易活动_上.png','图 4a  建立预约及待确认阶段的确认  取消 拒绝 屏蔽',width=9.55,maxheight=5.25)
para('本图描述一次预约。建立预约失败只结束本次请求；卖家只能确认待确认预约，普通取消恢复发布，屏蔽导致的取消保留商品已屏蔽状态。')
page('预约与交易活动图下半部分',True)
guide('导航 B  承接 A 的已确认预约，处理线下交付后的完成、取消或管理员屏蔽。')
pic(FIG/'交易活动_下.png','图 4b  已确认阶段的线下交付  买家完成 双方取消 屏蔽',width=9.55,maxheight=5.25)
para('买家只能完成已确认交易。取消后可主动新建预约，不能恢复旧预约；已完成记录保留首次完成时间。图中的屏蔽分支表示管理员屏蔽操作已经成功生效。')
page('买家预约商品时序图上半部分')
pic(FIG/'预约时序_上.png','图 5a  身份校验  锁定商品 重复请求识别及预约条件校验',maxheight=7.95)
para('参与者为买家、Web 界面、预约服务、用户服务、日志服务和数据库。用户身份从服务端会话取得；请求标识用于识别重发，页面版本用于识别陈旧请求。先查询成功请求，再判断是否建立新的预约。')
page('买家预约商品时序图下半部分')
guide('承接图5a的条件校验；参与者与上半部分相同，续图顶部按生命线位置重复标注。')
pic(FIG/'预约时序_下_含参与者.png','图 5b  建立预约  联动写入 提交与异常回滚',maxheight=6.4)
para('预约、商品、请求记录和日志使用同一事务，COMMIT 成功后才返回预约成功；任一步写入或提交失败均处理回滚。相同请求返回原预约及其当前状态，不恢复已取消预约；相同标识但内容不同则拒绝，新操作使用新标识。商品锁与有效预约检查使竞争请求仅有一个成功。')

page('四 各模块界面设计')
para('界面采用白底、蓝色主要操作、灰色背景和清晰状态标签。普通用户共用商品大厅、我的商品、我的预约入口；管理员使用审核、全部商品、统计和日志入口，首个管理员额外显示管理员账号入口。相同权限和状态使用一致的按钮规则。')
table(['业务模块','主要界面及编号','设计重点'],[
 ['用户与权限','07—12 登录注册；59—67 管理员账号','统一登录；注册固定普通用户；首个管理员最终批准'],
 ['商品管理','01—06 公开浏览；13—22 我的商品和编辑；36—48 管理员审核处理','字段校验、单图可选、审核及状态限制、屏蔽后恢复草稿'],
 ['预约与交易','03—04 发起预约；23—35 买入卖出记录和交易处理','两次点击预约；按身份和状态开放操作；保留取消与完成历史'],
 ['交易统计','49—51 日期查询及异常状态','按北京时间和预约价格统计；日期错误不冒充已应用'],
 ['操作日志','52—58 列表、组合查询、只读详情','各筛选条件同时满足；分别显示商品与预约前后状态']
],[1.1,2.75,3.25],10)
sub('1 原型与源文件')
para('离线交互原型使用 HTML、CSS 和 JavaScript 制作，通过本地 Chrome 导出 67 张 PNG，并保留对应 SVG 和七份 draw.io 源文件。在 diagrams.net 32.6.9 中加载、原生导出并保存全部 67 页，核对对象、文字、样式及位置；商品列表的一处文字完成了编辑、保存和恢复测试。')
sub('2 通用交互规则')
para('字段错误显示在对应位置并保留其他输入。状态过期时保留已填原因，提示刷新后按新状态操作；页面不会自动把一次失败请求转换成另一种业务操作。已完成记录只读；不属于当前用户的记录不开放修改入口。')
para('后续各图选择主要页面说明；其余错误、空结果和确认状态均在 UI界面 目录按编号保留。样例账号、商品和联系方式均为虚构数据。')

ui_pages=[
 ('用户登录与注册','07','图 6  用户登录界面','同一入口登录普通账号与管理员；普通注册只创建 MEMBER。游客可先浏览商品，发起预约后登录并返回原商品；凭据错误统一提示，密码以掩码显示。相关状态见 07—12。'),
 ('商品浏览与预约入口','01','图 7  公开商品列表','按名称、固定分类和价格组合查询，按最近发布时间倒序分页展示。详情保留描述中的联系与交付说明；登录后“立即预约”和“确认预约”两次点击完成预约。相关状态见 01—06。'),
 ('我的商品与发布编辑','15','图 8  编辑商品草稿','保存允许不完整；提交审核检查名称、分类、正数且至多两位小数的价格和描述。图片最多一张且可不传。待审核先撤回，发布或下架后先退草稿再修改；有效预约和完成历史受到保护。相关状态见 13—22。'),
 ('我的预约与交易处理','23','图 9  我的预约买入记录','买入和卖出页按同一普通账号切换。卖家确认或拒绝待确认预约，双方可填写原因取消待确认或已确认预约，只有买家可完成已确认交易。旧记录保留价格、取消信息和首次完成时间。相关状态见 23—35。'),
 ('管理员审核与违规处理','38','图 10  管理员商品审核详情','管理员只读核对商品内容，通过后发布，驳回需说明原因并退草稿。屏蔽未完成商品时取消有效预约，恢复仅回草稿，旧预约不恢复；已完成商品只读。相关状态见 36—48。'),
 ('管理员交易统计','49','图 11  管理员交易统计','起止日期按北京时间包含两端，按首次完成时间统计，金额汇总预约价格快照。默认虚构样例为 6 笔、390.50 元；空区间显示 0。日期错误保留旧结果并注明其有效范围。相关状态见 49—51。'),
 ('管理员操作日志','52','图 12  管理员操作日志','时间、操作人、商品编号、预约编号、类型按 AND 组合筛选。日志只读，分别保留商品与预约前后状态；仅涉及商品时不伪造预约信息，不记录密码。相关状态见 52—58。'),
 ('首个管理员批准并设置账号','60','图 13  批准并设置管理员账号','首个管理员填写新账号并最终确认批准，身份固定 ADMIN，不提供普通用户升级或线上申请队列。其他管理员仍可处理业务，但不能增设管理员；新账号可在当前演示会话内统一登录。相关状态见 59—67。')
]
ui_reasons={'07': '设计理由：统一入口减少身份选择错误，角色由账号权限确定；注册不提供管理员选项，避免把权限当作用户可填写字段。', '01': '设计理由：组合筛选和返回保留条件便于比较商品；预约前再次确认价格，在满足操作步数要求的同时减少误点。', '15': '设计理由：保存草稿和提交审核分开，支持分次填写；提交时集中校验必填项，减少无效商品进入审核队列。', '23': '设计理由：按身份和状态开放操作，区分卖家确认与买家完成；取消原因和处理历史帮助双方理解预约为什么终止。', '38': '设计理由：详情只读使管理员专注审核，不替卖家改内容；恢复后先回草稿，保留再次审核这一发布条件。', '49': '设计理由：使用预约价格快照保持历史金额稳定；明确标注预约价格合计，避免被理解为已经收取的货款。', '52': '设计理由：多个条件同时生效便于定位一次操作；商品与预约状态分开展示，可以检查屏蔽和取消是否联动。', '60': '设计理由：固定管理员身份并增加最终确认，减少误建账号；首个管理员按预置编号识别，权限不由界面输入决定。'}
for title,id,caption,text in ui_pages:
 page(title,True)
 if id=='07':
  # 与其余页面一致地从矢量源导出裁切图；没有修改正式界面。
  f=FIG/'界面_07.png'
 else:f=FIG/('界面_'+id+'.png')
 pic(f,caption,width=8.9,maxheight=5.5)
 p=para(text);p.paragraph_format.line_spacing=Pt(13.5);p.paragraph_format.space_after=Pt(0)
 for r in p.runs:r.font.size=Pt(10)
 p=para(ui_reasons[id]);p.paragraph_format.line_spacing=Pt(13.5);p.paragraph_format.space_after=Pt(0)
 for r in p.runs:r.font.size=Pt(10)

page('关键异常界面与设计理由',True)
para('下面两处提示把业务规则直接呈现在操作发生的位置，既说明当前操作为何不能继续，也给出下一步可执行的处理方式。')
t=table(['取消原因缺失','恢复已屏蔽商品'],[['',''],['图14a  保留输入并定位缺失原因','图14b  明确恢复结果为草稿'],['原因必填用于保留双方可理解的处理记录。错误提示紧邻输入框，保留其余信息，填写后才允许提交取消。','恢复不等于重新发布。确认框明确商品回到草稿、旧预约仍取消，使用户在操作前理解结果和重新审核要求。']],[4.75,4.75],10)
cellpic(t.rows[1].cells[0],FIG/'异常_取消原因.png',4.0,4.6)
cellpic(t.rows[1].cells[1],FIG/'异常_恢复草稿.png',4.0,4.6)

page('五 验证结果与实现边界')
sub('1 已执行的检查')
para('本地 PlantUML 检查通过 5 个 .puml 文件；类图、组件图、活动图、时序图共四种类型均有对应 PNG。67 张界面 PNG 均可读取，67 份 SVG 可解析；七份 draw.io 文件的页面和图片编号对应。')
para('整体核对在隔离副本重跑六组既有交互检查，共 229 项通过；追加 21 项跨图、跨页面和边界检查，共 250 项前端检查通过。检查覆盖字段校验、权限入口、状态流转、价格与日期边界、重复请求演示及跨页导航。')
sub('2 发现并修正的问题')
para('整体核对复现了连续预约两件商品时只保存第一条预约的问题：第二件商品已显示被预约，但成功页仍指向第一条记录。修订后按商品保存本次各条预约，使用不同编号，进入预约管理时一并带入并定位所选记录；取消一条不影响另一条。此外，无图商品改为明确的占位，避免将可选图片误显示为样例图片。')
para('类图中的商品状态为 8 种、预约状态为 4 种，关键操作类型为 12 种，与页面规则核对一致。新增管理员能进入业务后台，但不能继续增设管理员；首个管理员以固定编号校验。')
sub('3 验证范围与后续实现')
para('当前原型没有真实服务端和共享业务数据库。商品、预约、统计、日志主要使用各页内存中的独立样例；账号在同一标签页使用 sessionStorage 保存虚构的带盐校验信息。页面角色和 URL 参数属于演示上下文，不能替代真实身份凭证。')
para('正式实现需由服务端会话确定操作者，校验角色、对象归属和状态，数据库提供外键、唯一约束、事务、版本控制及持久请求记录。服务重启后的数据恢复、并发唯一预约、故障回滚和真实安全验收留待后续实现。')
para('性能目标为准备 1000 条商品数据，10 个并发用户完成共 100 次列表查询，其中成功且耗时不超过 2 秒的至少 95 次；失败和超时不能从分母删除，并须记录机器和配置。当前未做真实压测，设计上采用数据库过滤、排序、分页与合适索引。')

page('需求与设计对应关系上半部分')
reqs=[
 ('001','注册登录及新增管理员','UserService 和首管校验；07—12、59—67'),
 ('002','角色和对象权限','各服务鉴权及关联；游客与权限拒绝状态'),
 ('003','商品草稿和单张可选图','Product 可空字段；13—15'),
 ('004','提交撤回审核与驳回','发布审核活动图；16—18、36—41'),
 ('005','下架编辑和状态限制','takeDown、returnToDraft；19—22'),
 ('006','组合查询与分页','公开、本人及管理员查询；01、13、37'),
 ('007','唯一有效预约及价格快照','预约时序图、reservedPrice；02—04'),
 ('008','卖家确认预约','confirm 及关联更新；24、26'),
 ('009','双方取消与卖家拒绝','cancel、取消人原因时间；27—28、31、35'),
 ('010','买家完成及不可逆历史','complete、completedAt；29—30'),
 ('011','屏蔽取消与恢复草稿','block、cancelForBlocking、restore；32、42—46'),
 ('012','本人商品与买入卖出历史','listOwnProducts、listAsBuyer/listAsSeller；13、23—25')
]
table(['需求','主要内容','设计和界面对应'],[('REQ-'+a,b,c) for a,b,c in reqs],[.85,2.3,3.95],10)
page('需求与设计对应关系下半部分')
reqs=[
 ('013','唯一编号与历史关联','实体编号、普通关联、价格快照；关系数据库设计'),
 ('014','完成笔数与预约价格合计','StatisticsService；北京时间日期范围；49—51'),
 ('015','成功关键操作日志及查询','OperationLog、LogService；52—58'),
 ('016','持久化重复请求识别','OperationRequestRecord；预约时序图先查重'),
 ('017','查询性能目标','服务端组合过滤、排序、分页和索引；真实压测待完成'),
 ('018','密码和受限操作鉴权','User.passwordHash；服务端身份、角色及归属检查'),
 ('019','持久化与联动一致性','共享关系库、同一事务、版本与故障回滚设计'),
 ('020','操作步数与输入保留','两次点击预约；错误和旧页面提示不清空其他输入'),
 ('021','实现与业务范围','同一 Web 应用；五模块；本地图片；线下交易'),
 ('022','描述中联系和交付说明','Product.description；公开与相关双方查看，无独立聊天')
]
table(['需求','主要内容','设计和界面对应'],[('REQ-'+a,b,c) for a,b,c in reqs],[.85,2.3,3.95],10)
usage=[
 ('01 需求归类','将 22 条需求分到五个模块，给出跨模块项。','有助于建立整体结构；跨模块约束容易漏归属。','保留五模块，明确权限、日志及一致性由多模块共同承担。'),
 ('02 功能与协作','列出各模块功能和配合顺序。','便于理解调用；管理后台容易被误列为新模块。','后台作为入口分组，审核放商品模块，统计与日志单列。'),
 ('03 类图分析','分析实体、属性、关联和服务。','能给出初稿；实体增加过多会使初学者难以核对。','按实际需求保留核心实体、服务、枚举及统计结果。'),
 ('04 类图代码','依据方案生成 PlantUML。','便于快速形成可修改代码；代码正确不等于业务正确。','逐项核对字段和多重性，完整签名保留在源码。'),
 ('05 类图文件与图片','保存 .puml 并本地生成图片。','可重复生成；首次图过宽，阅读吃力。','根据本人“图太宽”的反馈进入布局优化。'),
 ('06 类图布局','重新组织分区、文字和连线。','改善阅读；简单缩图会让文字更小。','按实体、服务、枚举分区，简化展示并保存历史稿。'),
 ('07 类图复查','分析类图还有哪些可改进之处。','能发现约束不明确；过度设计会增加冗余。','本人要求简洁清晰，采用与需求直接相关的修订。'),
 ('08 三项类图修订','调整草稿方法、版本检查和首管权限。','强化关键边界；新增细节可能影响其他图。','区分新建和修改草稿，统一版本，明确首管校验。'),
 ('09 组件图分析','分析组件、依赖和数据存储。','能快速对应模块；依赖箭头易被读成时间顺序。','同一应用内五组件，图例说明依赖语义和事务边界。'),
 ('10 组件图生成','生成 .puml 与图片。','复用既有规则；跨层连线过多会拥挤。','用业务组合组件表示 HTTP 入口，数据范围在图例说明。'),
 ('11 活动图分析','分析发布、审核、预约和交易分支。','可整理主流程；单图容易混合多个生命周期。','先确定正常、取消、屏蔽和失败等不同出口。'),
 ('12 活动图初稿','生成初版活动图和图片。','便于整体检查；图过长，终止与成功容易混淆。','保留初稿，继续针对状态判断复查。'),
 ('13 活动图复查','指出混合流程和状态出口问题。','能辅助找错；未经审查的分支可能改变需求。','本人同意拆图并修正状态判断。'),
 ('14 活动图拆分','生成发布审核、预约交易两张图。','单次流程更清楚；跨图衔接仍需文字解释。','仅已发布作为发布成功，旧预约不自动恢复或重试。'),
 ('15 时序图设计与生成','形成买家预约时序图并优化。','补足调用与事务顺序；遗漏提交确认会夸大成功。','先查重，再校验新预约；日志同事务，提交后才成功。'),
 ('16 公开浏览原型','设计列表、详情、预约确认和异常页。','能快速看到交互；演示数据可能被当成真实业务。','明确虚构样例、两次点击预约及公开状态边界。'),
 ('17 登录注册原型','补登录、注册、错误与游客返回流程。','减少界面遗漏；前端身份不能证明安全。','普通注册固定 MEMBER，统一登录，正式鉴权留给服务端。'),
 ('18 商品管理原型','补我的商品及发布编辑等状态。','能检查字段和按钮；容易把草稿当作发布表单。','不完整可保存，提交严格校验，单张图片可选。'),
 ('19 预约管理原型','补买入卖出、确认、取消和完成。','角色操作可直接检查；状态与原因容易不一致。','按角色和当前状态开放操作，保留取消原因及历史。'),
 ('20 管理员商品原型','补审核、屏蔽与恢复。','覆盖管理员流程；恢复容易误做成重新发布。','屏蔽联动取消，恢复只回草稿，完成历史只读。'),
 ('21 统计日志原型','补统计、日志查询和只读详情。','便于核对结果；时区、价格和筛选组合易出错。','按北京时间、预约价格和 AND 条件检查边界。'),
 ('22 管理员账号原型','补首管批准及新账号统一登录。','能演示完整入口；演示会话容易被误认为持久系统。','固定首管编号，保存虚构带盐校验信息，说明会话边界。'),
 ('23 目录整理','清理旧图、中间数据和下载压缩包。','便于提交；清理不当可能丢源文件。','核验哈希，提交图与工作资料分开，保留设计历史。'),
 ('24 整体核对与修正','核对 22 需求、UML 和页面，执行前端检查。','发现连续预约记录问题；测试范围需准确表述。','修复多预约和无图占位，补非功能设计，保留 diff。'),
 ('25 下载参考资料核对','读取旧说明及两份往届示例。','可参考结构；旧说明与当前年度要求有差异。','以当前 PDF 为准，不采用旧截止日期和旧 AI 禁令。'),
 ('26 报告初稿编写','基于现有文件整理正文和 AI 过程。','能汇总材料；个人体会和完成状态可能被编造。','依据实际成果、工具和检查记录校对正文。'),
 ('27 工具核验与报告精简', '加载七组界面，验证文字编辑保存，建立 Git 记录并精简正文。', '仅能打开文件还不足以证明可编辑和可保存。', '原生核对67页，保留编辑与恢复副本；验证边界集中说明。'),
 ('28 报告案例与图文优化', '补三个修改案例、全貌导航、续图参与者及界面设计理由。', '通用优缺点不足以解释具体取舍，分区长图也需要全局定位。', '把采纳的建议写为讨论后的设计决定，保留AI发现问题及执行修改的来源。')
]
for batch in range(3):
 page('六 AI 使用过程' if batch==0 else 'AI 使用过程续表'+str(batch))
 if batch==0:
  para('本次使用 Codex 中的 AI 对话与本地文件工具辅助设计。本人确定项目和需求范围，逐步要求先分析、再生成代码或界面，并提出图太宽、简洁清晰、活动图拆分等反馈；源文件生成、导出和大部分技术修订由 AI 执行。')
  para('下表按实际设计过程整理各轮使用场景、AI 输出、作用与不足，以及修改结果。简单确认合并到对应轮次。')
 rows=usage[batch*9:(batch+1)*9] if batch<2 else usage[18:]
 table(['轮次与场景','AI 输出','作用与不足','修改优化后的结果'],rows,[1.15,1.85,1.95,2.15],9.5)
 if batch==2:
  para('设计演变证据包括五份历史 UML 源稿、修订差异、交互检查及原生编辑器保存记录。实验二在本地 lab2-design 分支建立当前基线 74b00e6，并提交 UI 工具核验记录 490e125；提交时间对应实际操作，提交说明标注 AI 执行范围。')

page('修改案例一 类图的简化与约束补充')
sub('使用场景与初稿问题')
para('在确定五个模块后，我让 AI 先分析类和关系，再生成 PlantUML。初稿集中展示较长的方法签名和多条关联，整体过宽。我明确反馈“这个图太宽了”，并提出“能做的简洁是最好，条理要清晰”。')
sub('我的取舍与修改结果')
para('经讨论后，我决定把简化重点放在表达方式上：实体、服务和枚举分区，图中缩短方法签名，完整参数保留在源码中。对 AI 提出的进一步建议，我采用了草稿方法、商品版本检查、首个管理员权限这三项直接影响业务边界的修改。')
table(['修改前','讨论后的取舍','修改后的设计'],[
 ['saveDraft 同时承担新建和修改','创建和修改的输入责任不同，应分别说明','createDraft 创建草稿；updateDraft 仅修改本人草稿'],
 ['商品写操作未统一体现版本参数','旧页面不能覆盖已生效的新状态','提交、撤回、审核、下架等统一接收 expectedVersion'],
 ['只有一般管理员校验','新增管理员权限应准确限定到首个管理员','initialAdminId 配置与 requireInitialAdmin 独立校验']
],[2.15,2.25,2.7],10)
sub('作用与验证')
para('AI 的作用是快速形成可修改的结构，并提示可能遗漏的边界；初稿的不足是展示密度过高，部分方法约束不够明确。我的选择是保留必要业务规则，减少重复表达。最终按图1a、图1b及状态表阅读，既能查看实体关系，也能核对服务职责。')
para('对照依据：设计历史中的原始宽版和三项调整前类图，以及 UML/类图.puml。后者明确保存两个草稿方法、商品版本参数和首个管理员校验，且通过本地 PlantUML 检查。')

page('修改案例二 活动图拆分与状态出口修正')
sub('使用场景与初稿问题')
para('初版把商品发布、管理员审核、买家预约和交易处理放在一张活动图中，混合多个生命周期。发布循环以“商品仍为草稿？”判断是否继续，容易把非草稿误读为已经发布；预约失败和取消后的回线也容易被理解为自动重试。')
sub('我的取舍与修改结果')
para('我要求 AI 复查这张活动图。在它提出拆分和状态判断建议后，我决定按“商品发布与审核”“预约与交易”两个生命周期拆成两张，并要求修正状态判断。判断依据是每张图应能说明本次流程为何结束，而结束本次请求并不一定表示发布或交易成功。')
table(['观察位置','修改前的问题','修改后的规则'],[
 ['发布流程出口','通过“是否仍为草稿”退出循环后进入预约','只有已发布进入成功出口；已屏蔽、草稿和其他状态分别处理'],
 ['取消与管理员屏蔽','屏蔽主要位于补充说明，主线容易只看到普通取消','主线明确普通取消恢复发布，屏蔽取消则保持商品已屏蔽'],
 ['失败与重新预约','重复回线可能让读者误认为系统持续重试','失败结束本次请求；取消后主动发起新预约并建立新记录']
],[1.35,2.75,3.0],10)
sub('作用与验证')
para('AI 帮助列出分支并生成代码，我采纳拆分方案并以清晰的状态出口作为检查标准。最终两张图分别处理发布审核和一次预约，四幅分区放大图沿 A、B 衔接，屏蔽、取消、完成和请求失败均保留各自含义。')
para('对照依据：设计历史/活动图拆分前/活动图.puml，以及 UML 中两份最终活动图。最终 .puml 通过本地语法检查；本次按状态条件逐项核对，没有把失败终止写成交易完成。')

page('修改案例三 连续预约不同商品的记录修正')
sub('使用场景与问题复现')
para('在我要求整体核对后，AI 按“同一买家先预约 p1，再预约 p2”的步骤复现了记录错误。浏览原型只保留一条全局预约：第二件商品已经变成 RESERVED，但保存的预约仍指向 p1，页面进入不可用状态。这个问题由核对操作发现，不是仅从截图推测。')
sub('我的取舍与修改结果')
para('依据复现结果，我采用“按商品分别保留本次预约，并在进入我的预约时一并带入”的修订方案。我的判断是商品占用状态与预约记录必须相互对应：预约第二件商品不应覆盖第一条，取消一条也不应释放另一件商品。文件修改及验证由 AI 执行。')
table(['核对项','修改前','修改后'],[
 ['预约记录组织','只保存一个 state.booking','以商品编号为键保存 state.bookings，并分别生成预约编号'],
 ['查看预约入口','只传递单条预约，可能指向错误商品','带入本次各条预约并定位当前所选记录'],
 ['取消后的影响','无法可靠表达两条独立的本次记录','取消其中一条，另一商品的预约保持有效']
],[1.3,2.6,3.2],10)
sub('作用与验证')
para('这个案例说明，AI 生成的界面即使单次操作正常，也可能遗漏连续操作。我的取舍是按每件商品保存本次预约，使成功提示、记录入口和商品占用状态相互对应，并保留各次预约的历史。')
para('已保存的检查确认：连续预约形成两条独立记录；第二件商品显示正确详情；第一件商品仍可查看预约；进入预约管理保留两条记录并正确定位；取消其中一条不影响另一条。这些检查属于离线原型，真实数据库一致性仍按第五章的实现边界验收。')
para('对照依据：整体核对/修正前复现.json、本轮修订.diff，以及跨图与原型核验结果.json。')

page('AI 辅助设计的思考与协作流程')
reflection_start=len(doc.paragraphs)
sub('1 本次过程反映的作用与局限')
para('AI 适合把明确需求转换为候选模块、UML 初稿、重复界面状态和可执行检查，也能降低修改 PlantUML 语法、导出图片和整理材料的工作量。先看分析再生成文件，使错误能在改动成本较低时被指出。')
para('AI 的初稿仍需要用业务约束检查。图形能够生成，并不保证状态判断正确；界面能够点击，也不保证事务和真实权限已经实现。类图过宽、活动图混合多个流程、连续预约只保留一条记录，都是本次需要继续复查和修订的具体问题。')
para('本次明确采用了“简洁、条理清晰”的判断标准。简化应通过明确职责、分区布局、拆分生命周期和去掉重复表达完成；同一商品唯一有效预约、取消原因、首个管理员权限等关键约束仍需保留。')
sub('2 可重复使用的协作流程')
for t in [
 '先固定需求编号、业务范围和必须保持的约束，避免每一轮自由扩展功能。',
 '先让 AI 分析模块、对象和分支，由本人检查是否符合需求，再生成源码和图片。',
 '按正常路径、权限边界、状态冲突、重复请求和失败回滚逐项审查，并要求给出可复现证据。',
 '一次聚焦一类改动，保留原稿、差异和验证结果；跨图或跨页面规则变动后同步核对。',
 '用 Git 保存每次修改，并在提交说明中区分设计判断、AI 生成和实际执行的改动。',
 '最后依据成果写报告，检查每项完成声明；本人复核正文、个人体会、图片和最终提交包。'
]:doc.add_paragraph(t,style='List Number')
sub('资料依据')
para('软件工程 2026 秋 实验二当前手册，2026 年 10 月 30 日下载的 PDF；实验一 requirements.md，Aulricy/SE_project 仓库 lab1-final 标签；本地最终 UML、界面源文件、UI设计说明与整体核对记录。两份往届示例仅参考文件组织方式，不复制其项目设计或个人经历。')
para('需求版本链接：https://github.com/Aulricy/SE_project/blob/lab1-final/requirements.md')

for p in doc.paragraphs[reflection_start:]:
 if not p.style.name.startswith('Heading'):
  p.paragraph_format.line_spacing=Pt(14.5);p.paragraph_format.space_after=Pt(5)
  for r in p.runs:r.font.size=Pt(10.5)
# 统一校正字号、保护孤行并清除 Word 标题主题色。
for p in doc.paragraphs:
 p.paragraph_format.widow_control=True
 if p.style.name.startswith('Heading') or p.style.name in ('Title','Subtitle'):
  for r in p.runs:r.font.color.rgb=RGBColor(0,0,0)
# 类图分区保留可放大的高分辨率图，不修改正式 UML。
doc.save(OUT)
(WORK/'报告页面计划.json').write_text(json.dumps(page_manifest,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'docx':str(OUT),'plannedPages':len(page_manifest),'images':len(doc.inline_shapes),'usageRounds':len(usage)},ensure_ascii=False))
